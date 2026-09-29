"""Document extractor abstractions.

Defines the common interface (``Extractor``) for turning a source document
into the ``(full_text, word_page_map)`` pair :func:`ingestion.chunker.chunk_document`
expects, following the Strategy / Driver pattern described in AGENTS.md.

Unlike every other Strategy in this project (``EMBEDDING_DRIVER``,
``CHUNKING_STRATEGY``, ...), the active extractor is **not** chosen from
``settings`` — it's chosen by the file's extension, via :func:`get_extractor`:
    - ``.pdf``                → :class:`PDFExtractor`
    - ``.md`` / ``.markdown`` → :class:`MarkdownExtractor`
    - ``.docx``               → :class:`DocxExtractor`

That's deliberate: which extractor applies is a fact about the file, not a
preference — there's nothing to configure.

Key exports:
    Extractor            -- Abstract base class defining the two-step extract contract.
    PDFExtractor         -- Concrete extractor for .pdf files.
    MarkdownExtractor    -- Concrete extractor for .md / .markdown files.
    DocxExtractor        -- Concrete extractor for .docx files.
    EXTRACTOR_REGISTRY   -- Dict mapping file extensions to their Extractor classes.
    SUPPORTED_EXTENSIONS -- frozenset of all registered extensions (derived from registry).
    get_extractor        -- Returns the correct Extractor instance for a given file path.
    normalize_extensions -- Normalises a collection of extension strings (adds dot, lowercases).

Usage::

    from ingestion.extractors import get_extractor
    extractor = get_extractor(Path("notes.md"))
    extractor.validate(path)
    full_text, word_page_map = extractor.extract(path)
"""

import zipfile
from abc import ABC, abstractmethod
from collections.abc import Iterable
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.exceptions import PackageNotFoundError

from ingestion.pdf_loader import extract_document_text, extract_pages, is_scanned_pdf


class Extractor(ABC):
    """Abstract base class for all document-format extractors.

    Two-step contract, mirroring how ``ingestion.ingest.add_document()``
    already used PDF extraction before this abstraction existed: validate
    first (fail with a clear error before any embedding-driver/DB work
    happens), then extract the real content.
    """

    @abstractmethod
    def validate(self, file_path: Path) -> None:
        """Raise if ``file_path`` has no usable content.

        Args:
            file_path: Path to the source document.

        Raises:
            FileNotFoundError: If the file doesn't exist.
            ValueError: If the file exists but has no extractable content
                (e.g. empty, or a scanned PDF with no text layer).
        """

    @abstractmethod
    def extract(self, file_path: Path, mode: str = "flat") -> tuple[str, list[int]]:
        """Return the whole document as one string, plus a per-word page/section map.

        Args:
            file_path: Path to the source document.
            mode: Forwarded to :func:`ingestion.pdf_loader.extract_document_text`
                for :class:`PDFExtractor` (``"flat"`` or ``"blocks"``).

        Returns:
            A ``(full_text, word_page_map)`` tuple.
        """

    def extract_with_headers(
        self, file_path: Path, mode: str = "flat"
    ) -> tuple[str, list[int], list[str] | None]:
        """Return full_text, word_page_map, and optional word_header_map.

        Extractors that do not extract headers (e.g. PDFExtractor) return
        ``None`` for ``word_header_map``.
        """
        full_text, word_page_map = self.extract(file_path, mode=mode)
        return full_text, word_page_map, None


class PDFExtractor(Extractor):
    """Wraps the existing, unchanged ``ingestion.pdf_loader`` functions.

    No PDF-parsing logic lives here — this class only adapts
    ``pdf_loader``'s functions to the :class:`Extractor` interface, reusing
    them exactly as ``ingestion.ingest.add_document()`` used to call them
    directly.
    """

    def validate(self, file_path: Path) -> None:
        pages = extract_pages(file_path)
        if not pages:
            raise ValueError(
                f"'{file_path.name}' has no pages. The file may be empty or corrupted."
            )
        if is_scanned_pdf(pages):
            raise ValueError(
                f"'{file_path.name}' appears to be a scanned PDF with no text layer. "
                "OCR support is not implemented in this version."
            )

    def extract(self, file_path: Path, mode: str = "flat") -> tuple[str, list[int]]:
        return extract_document_text(file_path, mode=mode)


class MarkdownExtractor(Extractor):
    """Reads a Markdown file directly — no coordinate-based heuristics needed.

    Unlike a PDF, Markdown already marks its own paragraph breaks (blank
    lines) and structure (ATX headers, ``# `` through ``###### ``) directly
    in the text — there's nothing to infer, so ``extract()`` just reads the
    file as-is. The ``mode`` parameter (PDF's flat-vs-blocks choice) doesn't
    apply here and is ignored, the same way :class:`ingestion.chunker.WordChunkingStrategy`
    ignores the ``driver`` parameter it's still handed for interface
    uniformity.
    """

    def validate(self, file_path: Path) -> None:
        if not file_path.exists():
            raise FileNotFoundError(f"Markdown file not found: {file_path}")
        if not file_path.read_text(encoding="utf-8").strip():
            raise ValueError(f"'{file_path.name}' is empty.")

    def extract(self, file_path: Path, mode: str = "flat") -> tuple[str, list[int]]:
        full_text = file_path.read_text(encoding="utf-8")
        word_page_map, _ = _markdown_headers_and_sections(full_text)
        return full_text, word_page_map

    def extract_with_headers(
        self, file_path: Path, mode: str = "flat"
    ) -> tuple[str, list[int], list[str] | None]:
        full_text = file_path.read_text(encoding="utf-8")
        word_page_map, word_header_map = _markdown_headers_and_sections(full_text)
        return full_text, word_page_map, word_header_map


def _markdown_headers_and_sections(full_text: str) -> tuple[list[int], list[str]]:
    """Parse Markdown lines, returning parallel section and header-breadcrumb maps.

    Tracks a stack of ATX headers (``#`` through ``######``) outside fenced code
    blocks. For each word in ``full_text.split()``, returns:
    1. A 1-based section counter (incremented on each header).
    2. A hierarchical breadcrumb string, e.g.
       ``"# Chapter 1 > ## Section 1.1"`` (empty string before any header).

    Args:
        full_text: The complete Markdown document text.

    Returns:
        A tuple of (word_section_map, word_header_map), both having length
        matching ``len(full_text.split())``.
    """
    word_section_map: list[int] = []
    word_header_map: list[str] = []
    section = 1
    in_code_fence = False
    header_stack: list[tuple[int, str]] = []

    for line in full_text.splitlines():
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            in_code_fence = not in_code_fence
        elif not in_code_fence and stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            # Standard ATX headers: 1 to 6 hashes, followed by space or empty
            if 1 <= level <= 6 and (len(stripped) == level or stripped[level] == " "):
                header_text = stripped[level:].strip().rstrip("#").strip()
                section += 1
                while header_stack and header_stack[-1][0] >= level:
                    header_stack.pop()
                header_stack.append((level, f"{'#' * level} {header_text}"))

        current_breadcrumb = " > ".join(h[1] for h in header_stack)
        num_words = len(line.split())
        word_section_map.extend([section] * num_words)
        word_header_map.extend([current_breadcrumb] * num_words)

    return word_section_map, word_header_map


def _markdown_section_map(full_text: str) -> list[int]:
    """Return a 1-based section index per word in ``full_text.split()``.

    Backward-compatible convenience wrapper around :func:`_markdown_headers_and_sections`.
    """
    word_section_map, _ = _markdown_headers_and_sections(full_text)
    return word_section_map


class DocxExtractor(Extractor):
    """Reads a .docx file, using paragraph **alignment** as the header signal.

    Word documents in this project's corpus (Hungarian court decisions)
    don't use Word's "Heading 1"/"Heading 2" paragraph styles for section
    titles — confirmed empirically across several real documents: every
    style-info lookup came back "Normal", never a heading style. What they
    *do* consistently use is **centered alignment** for section titles
    ("ítélete", "Indokolás", ...), confirmed the same way: centered
    paragraphs are rare (2-4 out of 100-400 per document) and are always a
    section title, never body text — a cheap, reliable, markup-level
    signal, unlike PDF's coordinate/font-size heuristics for the same job.
    Letter-spacing ("Í T É L E T") turned out not to be a reliable marker on
    its own (some titles are plain "INDOKOLÁS"), so alignment alone is what
    this checks.

    Unlike Markdown's ``#``...``######`` levels, these titles are flat, not
    hierarchical — there's no nesting to track, just "which title came most
    recently before this word", the same shape :func:`_markdown_headers_and_sections`
    produces, so it reuses the same ``ChunkMetadata.header_path``/section-index
    convention downstream.

    The ``mode`` parameter (PDF's flat-vs-blocks choice) doesn't apply here
    and is ignored, the same way :class:`MarkdownExtractor` ignores it.
    """

    def validate(self, file_path: Path) -> None:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")
        try:
            document = docx.Document(file_path)
        except (PackageNotFoundError, zipfile.BadZipFile) as exc:
            raise ValueError(f"'{file_path.name}' is not a valid DOCX file.") from exc
        if not any(p.text.strip() for p in document.paragraphs):
            raise ValueError(f"'{file_path.name}' has no extractable text content.")

    def extract(self, file_path: Path, mode: str = "flat") -> tuple[str, list[int]]:
        full_text, word_section_map, _ = _docx_text_and_headers(file_path)
        return full_text, word_section_map

    def extract_with_headers(
        self, file_path: Path, mode: str = "flat"
    ) -> tuple[str, list[int], list[str] | None]:
        return _docx_text_and_headers(file_path)


def _docx_text_and_headers(file_path: Path) -> tuple[str, list[int], list[str]]:
    """Read a .docx file's paragraphs into (full_text, word_section_map, word_header_map).

    Mirrors :func:`_markdown_headers_and_sections`'s output shape: a
    1-based section index and the most recent section title, per word in
    ``full_text.split()``. A paragraph becomes the new "current title" for
    every word from itself onward whenever it's center-aligned (see
    :class:`DocxExtractor`'s docstring for why that's the header signal here).

    Args:
        file_path: Path to the source .docx file.

    Returns:
        A ``(full_text, word_section_map, word_header_map)`` tuple.
    """
    document = docx.Document(file_path)

    full_text_parts: list[str] = []
    word_section_map: list[int] = []
    word_header_map: list[str] = []
    section = 1
    current_header = ""

    for paragraph in document.paragraphs:
        # Collapse embedded line breaks (python-docx renders <w:br/> as "\n"
        # inside .text) and stray runs of whitespace into single spaces.
        text = " ".join(paragraph.text.split())
        if not text:
            continue

        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            section += 1
            current_header = text

        full_text_parts.append(text)
        num_words = len(text.split())
        word_section_map.extend([section] * num_words)
        word_header_map.extend([current_header] * num_words)

    full_text = " ".join(full_text_parts)
    return full_text, word_section_map, word_header_map


#: Registry mapping lowercase file extensions to their Extractor classes.
EXTRACTOR_REGISTRY: dict[str, type[Extractor]] = {
    ".pdf": PDFExtractor,
    ".md": MarkdownExtractor,
    ".markdown": MarkdownExtractor,
    ".docx": DocxExtractor,
}

#: File extensions recognized by the ingestion pipeline.
SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(EXTRACTOR_REGISTRY.keys())


def normalize_extensions(extensions: Iterable[str]) -> frozenset[str]:
    """Normalize an iterable of file extensions to lowercase with leading dots.

    Args:
        extensions: An iterable of extension strings (e.g. ``[".md", "PDF", " .markdown "]``).

    Returns:
        A frozenset of normalized lowercase extensions with leading dots (e.g. ``{".md", ".pdf", ".markdown"}``).
    """
    return frozenset(
        ext.strip().lower()
        if ext.strip().startswith(".")
        else f".{ext.strip().lower()}"
        for ext in extensions
        if ext and ext.strip()
    )


def get_extractor(file_path: Path) -> Extractor:
    """Factory function: return the extractor matching ``file_path``'s extension.

    Args:
        file_path: Path to the source document.

    Returns:
        An :class:`Extractor` instance for that file's format.

    Raises:
        ValueError: If the file's extension isn't a supported format.
    """
    suffix = file_path.suffix.lower()
    extractor_cls = EXTRACTOR_REGISTRY.get(suffix)
    if extractor_cls is not None:
        return extractor_cls()

    supported_list = ", ".join(sorted(SUPPORTED_EXTENSIONS))
    raise ValueError(
        f"Unsupported document type: '{suffix}'. Supported extensions are: "
        f"{supported_list}."
    )
